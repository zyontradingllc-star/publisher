"""Build apps/platform/src/lib/blog/posts-generated.ts from xlsx + docx."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
from docx import Document

ROOT = Path(r"c:\Users\PC\Desktop\projects\next-js\book")
desktop = Path(r"c:\Users\PC\Desktop")
xlsx = next(desktop.rglob("50-BLOGS-With-Strong-Interlinks*.xlsx"))
docx_path = next(desktop.rglob("50-BLOGS-FULL-CONTENT-Ready*.docx"))
out_ts = ROOT / "apps/platform/src/lib/blog/posts-generated.ts"
out_json = ROOT / "apps/platform/src/lib/blog/generated-50.json"

IMAGES = [
    "/banners/publish-home-halloween.jpg",
    "/banners/home-banner.webp",
    "/banners/submit-book-buy-link-banner.webp",
    "/featured-row-2/verity.webp",
    "/featured-row-2/correspondent.webp",
    "/featured-row-2/the-hollow.webp",
    "/featured-row-2/shadow-of-eight.webp",
    "/featured-row-2/night-of-the-crash.webp",
    "/kobo/book-1.webp",
    "/kobo/book-2.webp",
    "/kobo/book-3.webp",
    "/kobo/book-4.webp",
    "/kobo/book-5.webp",
]


def infer_category(title: str, keyword: str, slug: str) -> str:
    blob = f"{title} {keyword} {slug}".lower()
    if any(x in blob for x in ("books2read", "booklinker", "vs ", "alternative", "compare")):
        return "Comparisons"
    if any(x in blob for x in ("amazon ads", "amazon ad", "kdp ads", "negative keyword")):
        return "Amazon Ads"
    if any(x in blob for x in ("newsletter", "swap", "email sequence")):
        return "Newsletter Swaps"
    if any(x in blob for x in ("launch", "preorder", "announce")):
        return "Book Launch"
    if any(x in blob for x in ("mindset", "imposter", "perfect", "productivity", "operating system", "deep work")):
        return "Author Mindset"
    if any(x in blob for x in ("royalt", "commission", "earn", "price", "money")):
        return "Royalties"
    if any(x in blob for x in ("list", "listing", "cover", "description", "genre", "kindle", "paperback", "audiobook")):
        return "Book Listing"
    if any(x in blob for x in ("bio", "link", "author page", "profile", "qr", "utm", "track", "click")):
        return "Author Page"
    if any(x in blob for x in ("instagram", "tiktok", "booktok", "facebook", "pinterest", "linkedin", "youtube", "social", "promote")):
        return "Promotion"
    return "Author Page"


def parse_h2s(raw: str) -> list[str]:
    headings = []
    for part in re.split(r"\s*\|\s*", str(raw or "")):
        part = part.strip()
        m = re.match(r"H2:\s*(.+?)(?:\s*\(|$)", part)
        if m:
            headings.append(m.group(1).strip())
        elif part:
            cleaned = re.sub(r"^H2:\s*", "", part).strip()
            if cleaned:
                headings.append(cleaned)
    headings = [h for h in headings if h][:3]
    while len(headings) < 3:
        headings.append(
            ["How to get started", "What to do next", "Why Authors Pilot"][len(headings)]
        )
    return headings


wb = openpyxl.load_workbook(xlsx, data_only=True)
ws = wb["50 Blogs - Ready to Publish"]
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

inter = wb["Interlink Map - Strong (5 Links Each)"]
inter_map: dict[str, list[dict[str, str]]] = {}
for r in range(2, inter.max_row + 1):
    slug = inter.cell(r, 2).value
    if not slug:
        continue
    links = []
    for c in range(3, 8):
        v = inter.cell(r, c).value
        if not v:
            continue
        href = str(v).strip()
        label = href.strip("/").split("/")[-1].replace("-", " ").title()
        links.append({"href": href if href.endswith("/") else href + "/", "label": label})
    key = str(slug).strip()
    if not key.endswith("/"):
        key += "/"
    inter_map[key] = links

# Parse DOCX into blogs by "Blog N:" headings
doc = Document(str(docx_path))
blocks: list[tuple[str, str]] = []
for p in doc.paragraphs:
    t = (p.text or "").strip()
    if not t:
        continue
    style = p.style.name if p.style else "Normal"
    blocks.append((style, t))

docx_by_num: dict[int, dict] = {}
current_n = None
current_sections: list[dict] = []
current_h3 = None
current_paras: list[str] = []


def flush_h3():
    global current_h3, current_paras
    if current_n is None:
        return
    if current_h3 and current_paras:
        current_sections.append({"heading": current_h3, "paragraphs": current_paras[:]})
    current_h3 = None
    current_paras = []


def flush_blog():
    global current_n, current_sections
    flush_h3()
    if current_n is not None:
        docx_by_num[current_n] = {"sections": current_sections[:]}
    current_sections = []


for style, text in blocks:
    m = re.match(r"Blog\s+(\d+)\s*:\s*(.+)$", text, re.I)
    if style.startswith("Heading 1") and m:
        flush_blog()
        current_n = int(m.group(1))
        continue
    if current_n is None:
        continue
    if style.startswith("Heading 3"):
        flush_h3()
        current_h3 = text
        continue
    if style.startswith("Heading 2"):
        # H1 duplicate — skip
        continue
    if text.lower().startswith("slug:") or text.lower().startswith("meta description:"):
        continue
    if style.startswith("List"):
        # keep related links as short lines under current section
        current_paras.append(text)
        continue
    current_paras.append(text)
flush_blog()

blogs = []
for r in range(2, ws.max_row + 1):
    row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
    title = str(row.get("Blog Title") or "").strip()
    if not title:
        continue
    no = int(row.get("No") or 0)
    slug_raw = str(row.get("Slug (/blog/...)") or "").strip()
    slug = slug_raw.replace("/blog/", "").strip("/")
    keyword = str(row.get("Primary Keyword") or "").strip()
    headings = parse_h2s(row.get("H2 Structure (AEO)"))
    aeo = str(row.get("AEO Direct Answer (45 words)") or "").strip()
    meta = str(row.get("Meta Description (155)") or "").strip()[:160]
    h1 = str(row.get("H1") or title).strip()
    category = infer_category(title, keyword, slug)
    slug_key = f"/blog/{slug}/"
    links = inter_map.get(slug_key, [])

    sections = docx_by_num.get(no, {}).get("sections", [])
    # Build body paragraphs: prefer step-by-step + answer + FAQ snippets
    body: list[str] = []
    if aeo:
        body.append(aeo)
    for sec in sections:
        h = sec["heading"].lower()
        if "related guides" in h or "call to action" in h:
            continue
        if "comparison table" in h:
            # skip table placeholder headings without useful prose
            paras = [p for p in sec["paragraphs"] if len(p) > 40]
            body.extend(paras[:1])
            continue
        if "faq" in h:
            # take first FAQ block as one paragraph
            if sec["paragraphs"]:
                body.append(sec["paragraphs"][0][:500])
            continue
        for p in sec["paragraphs"]:
            if len(p) < 30:
                continue
            if p.startswith("http") or p.startswith("1. http"):
                continue
            body.append(p)
            if len(body) >= 6:
                break
        if len(body) >= 6:
            break

    while len(body) < 3:
        body.append(
            f"{h1} helps indie authors share one branded Authors Pilot link, send readers to Amazon or their website, and keep 100% royalties with no commission."
        )

    # Align body with 3 H2s for article page (first 3 meaty paras)
    body_for_h2 = body[:3]
    if len(body) > 3:
        # keep extra in body for potential use; article maps headings[i] -> body[i]
        pass

    published = f"2026-{(no - 1) // 28 + 5:02d}-{((no - 1) % 28) + 1:02d}"
    # Prefer publish date from sheet if present
    pd = row.get("Publish Date (Daily 1)")
    if pd:
        try:
            published = pd.strftime("%Y-%m-%d") if hasattr(pd, "strftime") else str(pd)[:10]
        except Exception:
            pass

    blogs.append(
        {
            "slug": slug,
            "title": title,
            "h1": h1,
            "metaDescription": meta or f"Learn {keyword} on Authors Pilot. Free author page, one shareable link, 100% royalties.",
            "headings": headings,
            "excerpt": (aeo[:180] + "…") if len(aeo) > 180 else (aeo or meta[:180]),
            "category": category,
            "author": "Authors Pilot",
            "authorInitials": "AP",
            "publishedAt": published,
            "readTimeMinutes": max(5, min(12, 4 + len(" ".join(body)) // 900)),
            "image": IMAGES[(no - 1) % len(IMAGES)],
            "imageAlt": f"{h1} — Authors Pilot blog",
            "trending": no <= 3,
            "primaryKeyword": keyword,
            "body": body_for_h2 if len(body_for_h2) >= 3 else body[:3],
            "bodyExtra": body[3:6],
            "interlinks": links,
            "aeoAnswer": aeo,
        }
    )

out_json.write_text(json.dumps(blogs, indent=2, ensure_ascii=False), encoding="utf-8")
print(f"JSON {len(blogs)} posts -> {out_json}")

# Emit TypeScript
cats = sorted({b["category"] for b in blogs})


def esc(s: str) -> str:
    return json.dumps(s, ensure_ascii=False)


lines = []
lines.append("/** Auto-generated from 50-BLOGS xlsx + docx. Do not edit by hand. */")
lines.append('import type { BlogPost } from "./posts";')
lines.append("")
lines.append("export const GENERATED_BLOG_CATEGORIES = [")
for c in cats:
    lines.append(f"  {esc(c)},")
lines.append("] as const;")
lines.append("")
lines.append("export const GENERATED_BLOG_POSTS: BlogPost[] = [")
for b in blogs:
    lines.append("  {")
    lines.append(f"    slug: {esc(b['slug'])},")
    lines.append(f"    title: {esc(b['title'])},")
    lines.append(f"    h1: {esc(b['h1'])},")
    lines.append(f"    metaDescription: {esc(b['metaDescription'])},")
    lines.append(
        f"    headings: [{esc(b['headings'][0])}, {esc(b['headings'][1])}, {esc(b['headings'][2])}] as const,"
    )
    lines.append(f"    excerpt: {esc(b['excerpt'])},")
    lines.append(f"    category: {esc(b['category'])} as BlogPost['category'],")
    lines.append(f"    author: {esc(b['author'])},")
    lines.append(f"    authorInitials: {esc(b['authorInitials'])},")
    lines.append(f"    publishedAt: {esc(b['publishedAt'])},")
    lines.append(f"    readTimeMinutes: {b['readTimeMinutes']},")
    lines.append(f"    image: {esc(b['image'])},")
    lines.append(f"    imageAlt: {esc(b['imageAlt'])},")
    if b["trending"]:
        lines.append("    trending: true,")
    lines.append("    body: [")
    # include up to 6 paragraphs
    full_body = b["body"] + b.get("bodyExtra", [])
    for p in full_body[:6]:
        lines.append(f"      {esc(p)},")
    lines.append("    ],")
    lines.append(f"    interlinks: {json.dumps(b['interlinks'], ensure_ascii=False)},")
    lines.append(f"    aeoAnswer: {esc(b['aeoAnswer'])},")
    lines.append("  },")
lines.append("];")
lines.append("")

out_ts.write_text("\n".join(lines), encoding="utf-8")
print(f"TS -> {out_ts} ({out_ts.stat().st_size} bytes)")
print("Categories:", cats)
