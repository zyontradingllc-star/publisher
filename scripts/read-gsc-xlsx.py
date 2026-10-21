import openpyxl
from pathlib import Path

root = next(Path(r"C:\Users\PC\Desktop").glob("Dezyon Studio*"))
book = root / "book"
files = [
    "authorspilot.com-Coverage-2026-10-20.xlsx",
    "authorspilot.com-Coverage-Drilldown-2026-10-20.xlsx",
    "authorspilot.com-Coverage-Valid-2026-10-20.xlsx",
    "authorspilot.com-Coverage-Valid-2026-10-20 (1).xlsx",
    "authorspilot_seo_100_page_implementation_plan.xlsx",
]

out = Path(r"C:\Users\PC\Desktop\projects\next-js\book\scripts\gsc-xlsx-dump.txt")
lines = []

for name in files:
    path = book / name
    lines.append("\n" + "=" * 80)
    lines.append(name)
    lines.append("=" * 80)
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f"\n--- Sheet: {sheet} (rows={ws.max_row}, cols={ws.max_column}) ---")
        rows = list(ws.iter_rows(values_only=True))
        for i, row in enumerate(rows):
            cells = ["" if c is None else str(c).replace("\n", " ")[:120] for c in row]
            if not any(cells):
                continue
            lines.append(f"{i:03d}| " + " || ".join(cells))

out.write_text("\n".join(lines), encoding="utf-8")
print(f"wrote {out} ({len(lines)} lines)")
